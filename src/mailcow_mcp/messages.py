"""Reading messages: summaries, text bodies, attachment text extraction."""

from __future__ import annotations

import html
import io
import re
from datetime import UTC, date, datetime
from email.message import Message
from email.utils import getaddresses, parsedate_to_datetime

import nh3
from pydantic import BaseModel, Field

from mailcow_mcp.errors import InvalidInput
from mailcow_mcp.mime import is_attachment, part_bytes, safe_filename, walk_parts

SUMMARY_HEADERS = (
    "From",
    "To",
    "Cc",
    "Subject",
    "Date",
    "Message-ID",
    "In-Reply-To",
    "References",
    "Content-Type",
    "X-Spam-Score",
    "X-Rspamd-Score",
    "X-Spamd-Result",
    "X-Spam-Status",
)

MAX_BODY_CHARS = 100_000
MAX_EXTRACT_CHARS = 50_000
MAX_EXTRACT_BYTES = 5 * 1024 * 1024
MAX_PDF_PAGES = 200

_BLOCK_END = re.compile(
    r"</(p|div|tr|li|h[1-6]|blockquote|pre|table|section|article|header|footer)\s*>|<br\s*/?>|<hr\s*/?>",
    re.IGNORECASE,
)
_CELL_END = re.compile(r"</t[dh]\s*>", re.IGNORECASE)
_BLANK_LINES = re.compile(r"\n[ \t]*\n(?:[ \t]*\n)+")
_SPAM_RESULT = re.compile(r"\[\s*(-?\d+(?:\.\d+)?)\s*/")
_SPAM_STATUS = re.compile(r"score=(-?\d+(?:\.\d+)?)")


class MessageSummary(BaseModel):
    uid: int
    folder: str
    message_id: str | None = None
    date: str | None = Field(default=None, description="ISO 8601")
    sender: str | None = None
    to: list[str] = []
    cc: list[str] = []
    subject: str = ""
    seen: bool = False
    flagged: bool = False
    answered: bool = False
    size: int | None = None
    has_attachments: bool = False
    spam_score: float | None = None


class AttachmentInfo(BaseModel):
    part_id: str
    filename: str
    mime_type: str
    size: int


def header(message: Message, name: str) -> str:
    value = message.get(name)
    return " ".join(str(value).split()) if value is not None else ""


def addresses(message: Message, name: str) -> list[str]:
    values = [str(v) for v in message.get_all(name, [])]
    result = []
    for display, addr in getaddresses(values):
        if not addr:
            continue
        display = " ".join(display.split())
        result.append(f"{display} <{addr}>" if display else addr)
    return result


def message_date(message: Message, fallback: datetime | None = None) -> datetime | None:
    raw = message.get("Date")
    if raw:
        try:
            parsed = parsedate_to_datetime(str(raw))
            return parsed if parsed.tzinfo else parsed.replace(tzinfo=UTC)
        except (TypeError, ValueError, IndexError):
            pass
    if fallback is not None:
        return fallback if fallback.tzinfo else fallback.replace(tzinfo=UTC)
    return None


def spam_score(message: Message) -> float | None:
    for name in ("X-Rspamd-Score", "X-Spam-Score"):
        value = header(message, name)
        try:
            return float(value)
        except ValueError:
            pass
    for name, pattern in (("X-Spamd-Result", _SPAM_RESULT), ("X-Spam-Status", _SPAM_STATUS)):
        match = pattern.search(header(message, name))
        if match:
            return float(match.group(1))
    return None


def summarize(
    uid: int,
    folder: str,
    headers: Message,
    flags: tuple[bytes, ...],
    size: int | None = None,
    internal_date: datetime | None = None,
) -> MessageSummary:
    when = message_date(headers, internal_date)
    from_ = addresses(headers, "From")
    content_type = header(headers, "Content-Type").lower()
    return MessageSummary(
        uid=uid,
        folder=folder,
        message_id=header(headers, "Message-ID") or None,
        date=when.isoformat() if when else None,
        sender=from_[0] if from_ else None,
        to=addresses(headers, "To"),
        cc=addresses(headers, "Cc"),
        subject=header(headers, "Subject"),
        seen=b"\\Seen" in flags,
        flagged=b"\\Flagged" in flags,
        answered=b"\\Answered" in flags,
        size=size,
        has_attachments=content_type.startswith(("multipart/mixed", "multipart/report")),
        spam_score=spam_score(headers),
    )


def html_to_text(markup: str) -> str:
    """Readable text from HTML: block breaks kept, scripts and styles dropped."""
    markup = _BLOCK_END.sub(lambda m: m.group(0) + "\n", markup)
    markup = _CELL_END.sub(lambda m: m.group(0) + "\t", markup)
    text = nh3.clean(markup, tags=set(), clean_content_tags={"script", "style", "head", "title"})
    text = html.unescape(text).replace("\xa0", " ")
    lines = [line.rstrip() for line in text.splitlines()]
    return _BLANK_LINES.sub("\n\n", "\n".join(lines)).strip()


def _text_of(part: Message) -> str:
    try:
        content = part.get_content()  # type: ignore[attr-defined]
    except (LookupError, KeyError, ValueError, AssertionError):
        payload = part.get_payload(decode=True)
        content = payload.decode("utf-8", "replace") if isinstance(payload, bytes) else ""
    return content if isinstance(content, str) else ""


def body_text(message: Message) -> tuple[str, str]:
    """The message's text: (text, source) where source is "text/plain", "text/html" or ""."""
    plain = html_part = None
    for _, part in walk_parts(message):
        if is_attachment(part):
            continue
        if part.get_content_type() == "text/plain" and plain is None:
            plain = part
        elif part.get_content_type() == "text/html" and html_part is None:
            html_part = part
    if plain is not None:
        return _text_of(plain).replace("\r\n", "\n").strip(), "text/plain"
    if html_part is not None:
        return html_to_text(_text_of(html_part)), "text/html"
    return "", ""


def attachments(message: Message) -> list[AttachmentInfo]:
    result = []
    for number, part in walk_parts(message):
        if not is_attachment(part):
            continue
        result.append(
            AttachmentInfo(
                part_id=number,
                filename=safe_filename(part.get_filename(), f"part-{number}"),
                mime_type=part.get_content_type(),
                size=len(part_bytes(part)),
            )
        )
    return result


def iso_timestamp(value: str | None) -> float:
    """Sort key for ISO dates with any offset (unknown dates sort first)."""
    if not value:
        return 0.0
    try:
        parsed = datetime.fromisoformat(value)
    except ValueError:
        return 0.0
    return (parsed if parsed.tzinfo else parsed.replace(tzinfo=UTC)).timestamp()


def truncate(text: str, limit: int) -> tuple[str, bool]:
    if len(text) <= limit:
        return text, False
    return text[:limit] + "\n[… truncated]", True


# --- attachment text extraction ---------------------------------------------

TEXT_TYPES = {
    "application/json",
    "application/xml",
    "application/csv",
    "application/x-yaml",
    "application/yaml",
}
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
IMAGE_TYPES = {"image/png", "image/jpeg", "image/gif", "image/webp"}


class ExtractionError(Exception):
    pass


def extract_text(data: bytes, mime_type: str, filename: str, charset: str | None) -> str | None:
    """Text of a PDF, DOCX or text-like attachment; None if the type isn't supported."""
    name = filename.lower()
    if mime_type == "application/pdf" or name.endswith(".pdf"):
        return _pdf_text(data)
    if mime_type == DOCX or name.endswith(".docx"):
        return _docx_text(data)
    if (
        mime_type.startswith("text/")
        or mime_type in TEXT_TYPES
        or name.endswith(
            (".txt", ".csv", ".md", ".json", ".xml", ".log", ".ics", ".vcf", ".yaml", ".yml")
        )
    ):
        text = (
            data.decode(charset or "utf-8", "replace")
            if _known(charset)
            else data.decode("utf-8", "replace")
        )
        if mime_type == "text/html" or name.endswith((".html", ".htm")):
            return html_to_text(text)
        return text
    return None


def _known(charset: str | None) -> bool:
    if not charset:
        return False
    try:
        "".encode(charset)
    except LookupError:
        return False
    return True


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ExtractionError("The PDF is encrypted.")
        pages = []
        for page in reader.pages[:MAX_PDF_PAGES]:
            pages.append(page.extract_text() or "")
            if sum(len(p) for p in pages) > MAX_EXTRACT_CHARS:
                break
    except PdfReadError as exc:
        raise ExtractionError(f"The PDF can't be read: {exc}") from exc
    return "\n\n".join(p.strip() for p in pages).strip()


def _docx_text(data: bytes) -> str:
    import zipfile

    from docx import Document

    try:
        document = Document(io.BytesIO(data))
    except (zipfile.BadZipFile, KeyError, ValueError) as exc:
        raise ExtractionError("The DOCX file can't be read.") from exc
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(lines).strip()


def parse_day(value: str | None, name: str) -> date | None:
    if not value:
        return None
    try:
        return date.fromisoformat(value[:10])
    except ValueError as exc:
        raise InvalidInput(f"{name} must be a date like 2026-10-01.") from exc
